#!/usr/bin/env python3
"""Create Subtypes.xlsx defining which countries each subtype is relevant for."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Country assignment by type:
#   Spain  → e-commerce / retail focus
#   France → marketing / media / analytics focus
#   Italy  → enterprise IT / security focus
#   Shared across all: Account, Feedback, Service, Support, Transaction, User

COUNTRY_BY_TYPE = {
    'Account':       ['Spain', 'France', 'Italy'],
    'Alert':         ['Italy'],
    'Analytics':     ['France', 'Italy'],
    'Communication': ['France'],
    'Complaint':     ['Spain', 'Italy'],
    'Content':       ['France'],
    'Document':      ['Spain', 'Italy'],
    'Event':         ['Spain', 'France'],
    'Feedback':      ['Spain', 'France', 'Italy'],
    'Marketing':     ['France'],
    'Payment':       ['Spain'],
    'Product':       ['Spain', 'France'],
    'Return':        ['Spain'],
    'Security':      ['Italy'],
    'Service':       ['Spain', 'France', 'Italy'],
    'Shipping':      ['Spain'],
    'Support':       ['Spain', 'France', 'Italy'],
    'System':        ['Italy'],
    'Transaction':   ['Spain', 'France', 'Italy'],
    'User':          ['Spain', 'France', 'Italy'],
}

SUBTYPES_BY_TYPE = {
    'Account':       ['Developer', 'Education', 'Enterprise', 'Government', 'Nonprofit',
                      'Premium', 'Standard', 'Startup', 'Suspended', 'Trial', 'Unverified', 'Verified'],
    'Alert':         ['Debug', 'Error', 'Info', 'Maintenance', 'Notification', 'Performance',
                      'Reminder', 'Security', 'Status', 'Success', 'Urgent', 'Warning'],
    'Analytics':     ['Attribution', 'Click', 'Cohort', 'Conversion', 'Engagement', 'Funnel',
                      'Pageview', 'Retention', 'Revenue', 'Session', 'Traffic', 'User'],
    'Communication': ['Alert', 'Blog', 'Chat', 'Email', 'Forum', 'In App Message',
                      'Newsletter', 'Phone Call', 'Push Notification', 'SMS', 'Video Call', 'WhatsApp'],
    'Complaint':     ['Billing', 'Communication', 'Delivery', 'Discriminatory', 'Fraud',
                      'Misrepresentation', 'Pricing', 'Privacy', 'Quality', 'Safety', 'Service', 'Support'],
    'Content':       ['Archive', 'Audio', 'Code', 'Comment', 'Dataset', 'Document',
                      'Image', 'Presentation', 'Spreadsheet', 'Text', 'Video', 'Webpage'],
    'Document':      ['Certificate', 'Contract', 'Guide', 'Invoice', 'License', 'Manual',
                      'Policy', 'Receipt', 'Report', 'Shipping Label', 'Terms', 'Warranty'],
    'Event':         ['Conference', 'Exclusive', 'Hybrid', 'In Person', 'Launch', 'Meetup',
                      'Private', 'Public', 'Sponsored', 'Virtual', 'Webinar', 'Workshop'],
    'Feedback':      ['Bug Report', 'Comment', 'Feature Request', 'Improvement', 'Negative',
                      'Neutral', 'Positive', 'Question', 'Rating', 'Review', 'Suggestion', 'Testimonial'],
    'Marketing':     ['Affiliate', 'Content', 'Direct Mail', 'Email Campaign', 'Event',
                      'Influencer', 'Paid Ads', 'Partnership', 'PR', 'Referral', 'Social Media', 'Webinar'],
    'Payment':       ['Bank Transfer', 'Cash', 'Check', 'COD', 'Credit Card', 'Cryptocurrency',
                      'Debit Card', 'Digital Wallet', 'Gift Card', 'Invoice', 'Prepaid', 'Subscription'],
    'Product':       ['Bundle', 'Clearance', 'Custom', 'Digital', 'Exclusive', 'Gift',
                      'Licensed', 'Limited', 'Physical', 'Refurbished', 'Rental', 'Subscription'],
    'Return':        ['Changed Mind', 'Damaged in Transit', 'Defective', 'Duplicate', 'Expired',
                      'Full', 'Incompatible', 'Not As Described', 'Partial', 'Recall', 'Unwanted', 'Wrong Item'],
    'Security':      ['Access Control', 'Audit', 'Encryption', 'Firewall', 'Login', 'Password',
                      'Patch', 'Permission', 'Role', 'SSL', 'VPN', 'Vulnerability'],
    'Service':       ['Consultation', 'Delivery', 'Express', 'Installation', 'Maintenance',
                      'Onsite', 'Phone', 'Premium', 'Remote', 'Standard', 'Support', 'Training'],
    'Shipping':      ['Drop Off', 'Express', 'Fragile', 'Hazmat', 'Insured', 'International',
                      'Overnight', 'Self Pickup', 'Signature', 'Standard', 'Temperature Controlled', 'White Glove'],
    'Support':       ['Callback', 'Chat', 'Community', 'Email', 'Escalation', 'FAQ',
                      'In Person', 'Knowledge Base', 'Phone', 'Social Media', 'Ticket', 'Video Tutorial'],
    'System':        ['Authentication', 'Backup', 'Cache', 'Database', 'Load Balancer', 'Logging',
                      'Monitoring', 'Network', 'Queue', 'Security', 'Server', 'Storage'],
    'Transaction':   ['Bulk', 'Contract', 'Coupon', 'Discount', 'Exchange', 'Insurance',
                      'Promotion', 'Purchase', 'Refund', 'Subscription', 'Trial', 'Warranty'],
    'User':          ['Admin', 'Corporate', 'Customer', 'Free', 'Guest', 'Inactive',
                      'Individual', 'Influencer', 'Internal', 'Partner', 'Premium', 'Vendor'],
}

# --- Build rows ---
rows = []
for type_name in sorted(SUBTYPES_BY_TYPE):
    countries = COUNTRY_BY_TYPE[type_name]
    for subtype in SUBTYPES_BY_TYPE[type_name]:
        rows.append((subtype, type_name, countries))

# --- Create workbook ---
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Subtypes'

# Styles
header_font  = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
header_fill  = PatternFill('solid', fgColor='2E4057')
cell_font    = Font(name='Calibri', size=10)
alt_fill     = PatternFill('solid', fgColor='F0F4F8')
border_side  = Side(style='thin', color='CCCCCC')
cell_border  = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

country_colors = {'Spain': 'D4EDDA', 'France': 'CCE5FF', 'Italy': 'FFF3CD'}

# Headers
headers = ['Subtype', 'Type', 'Countries']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = cell_border

ws.row_dimensions[1].height = 22

# Data rows
for row_idx, (subtype, type_name, countries) in enumerate(rows, 2):
    fill = alt_fill if row_idx % 2 == 0 else None

    c1 = ws.cell(row=row_idx, column=1, value=subtype)
    c2 = ws.cell(row=row_idx, column=2, value=type_name)
    countries_str = '[' + ', '.join(f"'{c}'" for c in countries) + ']'
    c3 = ws.cell(row=row_idx, column=3, value=countries_str)

    for c in (c1, c2, c3):
        c.font = cell_font
        c.border = cell_border
        c.alignment = Alignment(vertical='center')
        if fill:
            c.fill = fill

    # Colour the countries cell by country membership
    if len(countries) == 1:
        c3.fill = PatternFill('solid', fgColor=country_colors[countries[0]])
    elif len(countries) == 2:
        # Light blend — use a light purple for 2-country overlap
        c3.fill = PatternFill('solid', fgColor='E8D5F5')
    # 3 countries → default (alternating) fill

ws.row_dimensions[row_idx].height = 16

# Column widths
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 30

# Freeze header row
ws.freeze_panes = 'A2'

# Auto-filter
ws.auto_filter.ref = f'A1:C{len(rows)+1}'

out_path = '/home/eyalshw/github/wzender/magellan/data/Subtypes.xlsx'
wb.save(out_path)
print(f'Saved {out_path} ({len(rows)} rows)')

# Print summary
from collections import Counter
by_country = Counter()
for _, _, countries in rows:
    for c in countries:
        by_country[c] += 1
for country, n in sorted(by_country.items()):
    print(f'  {country}: {n} subtypes')
